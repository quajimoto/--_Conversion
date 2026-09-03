from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI 경진대회 모바일 평가 시스템\n사용자 및 관리자 매뉴얼"
    subtitle.text = "부서별 평가 및 결과 통합 관리 시스템\n2026. 08."

    # 2. Overview Slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "1. 시스템 개요 및 주요 기능"
    tf = body_shape.text_frame
    tf.text = "본 시스템은 AI 경진대회의 부서별 평가를 모바일과 웹에서 쉽게 진행할 수 있도록 개발되었습니다."
    p = tf.add_paragraph()
    p.text = "평가자 (일반 사용자)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "간편한 로그인 (설정된 방식에 따라 리스트, ID/PW, 직접 입력 지원)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "부서별 직관적인 점수 입력 및 임시저장 기능"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "관리자 (Admin)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "실시간 평가 결과 통합 대시보드 조회 및 엑셀 다운로드"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "평가 내역 수정(삭제/복구) 및 평가 항목/부서 동적 관리"
    p.level = 2

    # 3. Login Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "2. 로그인 방법"
    tf = body_shape.text_frame
    tf.text = "관리자가 설정한 3가지 로그인 방식 중 하나로 접속됩니다."
    p = tf.add_paragraph()
    p.text = "리스트 선택 방식 (기본)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "이름 직접 입력 방식"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "ID / Password 방식"
    p.level = 1
    try:
        slide.shapes.add_picture('login.png', Inches(5.5), Inches(1.5), height=Inches(5))
    except Exception as e:
        pass

    # 4. Evaluator Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "3. 평가 진행 (평가자 모드)"
    tf = body_shape.text_frame
    tf.text = "직관적인 인터페이스로 빠르고 정확한 평가를 지원합니다."
    p = tf.add_paragraph()
    p.text = "부서 선택 및 점수 입력"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "저장 및 제출"
    p.level = 1
    try:
        slide.shapes.add_picture('eval.png', Inches(5.5), Inches(1.5), height=Inches(5))
    except Exception as e:
        pass

    # 5. Admin Dashboard Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "4. 결과 확인 및 다운로드 (관리자 모드)"
    tf = body_shape.text_frame
    tf.text = "모든 평가 데이터는 관리자 대시보드에서 실시간 통합됩니다."
    p = tf.add_paragraph()
    p.text = "결과 집계 탭 (선택한 날짜/부서별 결과)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "엑셀 다운로드 (현재 부서 / 전체 부서)"
    p.level = 1
    try:
        slide.shapes.add_picture('admin.png', Inches(4.5), Inches(1.5), height=Inches(5))
    except Exception as e:
        pass

    # 6. Admin Management Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "5. 평가 데이터 및 시스템 설정 (관리자 모드)"
    tf = body_shape.text_frame
    tf.text = "잘못된 평가를 수정하거나 시스템 환경을 커스터마이징 합니다."
    p = tf.add_paragraph()
    p.text = "평가 내역 수정 탭"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "잘못 제출된 평가 항목을 체크박스로 선택하여 '삭제 처리' 할 수 있습니다."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "실수로 삭제한 데이터는 '복구' 버튼을 통해 즉시 원상 복구됩니다."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "평가 항목 및 부서 관리 탭"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "부서를 추가/삭제하고 드래그 앤 드롭으로 표시 순서를 변경합니다."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "평가 항목명, 세부 설명, 최대 점수를 자유롭게 수정하고 항목을 추가할 수 있습니다."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "시스템 로그인 방식 변경 (우측 상단 톱니바퀴/드롭다운 '저장')"
    p.level = 1

    prs.save('c:/Users/user/Desktop/오정환/Antigravity/평가_Conversion/AI_Competition_Mobile_Web_Manual.pptx')

if __name__ == '__main__':
    create_presentation()
